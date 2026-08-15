from pathlib import Path
import shutil

from openpyxl import load_workbook


ROOT = Path("/Users/sunmengge/Dropbox/GitHub/sicence nationism")
WORKBOOKS = [
    ROOT / "code/Brexit-stata_external_try/brexit_non_eu.xlsx",
    ROOT / "code/crimes/original_completed.xlsx",
    ROOT / "code/crimes_2015_2022/original_completed.xlsx",
]


def normalize_country(value):
    if value is None:
        return value
    name = str(value).strip().upper()
    name = name.replace(".", "")
    name = name.replace(",", "")
    name = " ".join(name.split())
    replacements = {
        "KOREA REP": "KOREA",
        "COSTA RICA": "COSTA_RICA",
        "RUSSIAN FEDERATION": "RUSSIAN_FEDERATION",
        "SLOVAK REPUBLIC": "SLOVAK_REPUBLIC",
        "SOUTH AFRICA": "SOUTH_AFRICA",
        "UNITED KINGDOM": "UNITED_KINGDOM",
        "UNITED STATES": "UNITED_STATES",
        "NEW ZEALAND": "NEW_ZEALAND",
    }
    name = replacements.get(name, name)
    name = name.replace(" ", "_")
    return name


def normalize_workbook(path):
    backup = path.with_suffix(path.suffix + ".bak_country_names")
    if not backup.exists():
        shutil.copy2(path, backup)

    wb = load_workbook(path)
    changed = 0
    for ws in wb.worksheets:
        header = [cell.value for cell in ws[1]]
        if "Country" not in header:
            continue
        col = header.index("Country") + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            new_value = normalize_country(cell.value)
            if cell.value != new_value:
                cell.value = new_value
                changed += 1
    wb.save(path)
    return changed


if __name__ == "__main__":
    for workbook in WORKBOOKS:
        changed = normalize_workbook(workbook)
        print(f"{workbook}: normalized {changed} country cells")
